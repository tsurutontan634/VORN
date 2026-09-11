"""様式の書類出力。

LLM がコロニー情報と一次資料（様式・記載例）から「項目名 → 記入内容」を作る。
配布様式の DOCX があるときは、その表のセル構造（表番号・行・列・現在の文言）を LLM に渡し、
「どのセルにどう書くか」を LLM に決めさせる。コード側は指示どおりにセルを書き換えるだけ。
様式の項目名やセル位置はコードに固定しない。
"""
from __future__ import annotations

import datetime as dt
import io
from pathlib import Path

from docx import Document
from docx.shared import Pt

from .colony import Colony
from .profile import MunicipalityProfile

FIELDS_SCHEMA = {
    "type": "object",
    "additionalProperties": False,
    "properties": {
        "form_title": {"type": "string"},
        "addressed_to": {"type": "string"},
        "fields": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "label": {"type": "string"},
                    "value": {"type": "string"},
                    "status": {"type": "string", "enum": ["記入済", "要確認", "未入力"]},
                },
                "required": ["label", "value", "status"],
            },
        },
        "cell_edits": {
            "type": "array",
            "items": {
                "type": "object",
                "additionalProperties": False,
                "properties": {
                    "table": {"type": "integer"},
                    "row": {"type": "integer"},
                    "col": {"type": "integer"},
                    "text": {"type": "string"},
                },
                "required": ["table", "row", "col", "text"],
            },
        },
        "notes_for_applicant": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["form_title", "addressed_to", "fields", "cell_edits", "notes_for_applicant"],
}

FILL_SYSTEM = """あなたは自治体の地域猫活動支援制度の申請書類を代筆する事務担当者です。
一次資料の中の様式と記載例に沿って、コロニー情報から各記入欄を埋めてください。
- fields の label は様式にある項目名をそのまま使い、様式の並び順に出す。
- 記載例の文体（簡潔・事務的）に合わせる。情緒的な表現は使わない。
- コロニー情報に無い欄は value を空にし status を「未入力」にする。推測で埋めない。
- 情報はあるが記載例と照らして確認が要るものは「要確認」にし、理由を notes_for_applicant に書く。
- 日付は「令和◯年◯月◯日」形式に直す。
- notes_for_applicant は活動者に見せる文。JSON のキー名（members, cats など）は書かず、「活動者の氏名」「猫の一覧」のように日本語の項目名で書く。
- 様式テンプレートのセル一覧が渡されたときは、cell_edits に「そのセルの新しい全文」を出す。
  元のセルにある項目名や単位（「頭」「名」「か所」「電話　－」など）は消さずに残し、その中に値を入れる。
  チェック欄（□）は該当するものを ☑ に置き換える。書き換え不要なセルは cell_edits に含めない。
  セル一覧が無いときは cell_edits を空にする。"""


def template_cells(path: Path) -> list[dict]:
    """テンプレート DOCX の表セルを (table,row,col,text) で列挙する。結合セルは最初の位置だけ。"""
    doc = Document(str(path))
    out = []
    for ti, table in enumerate(doc.tables):
        seen: list = []  # 結合セルは同じ要素が複数回返るので、要素そのもので重複を除く
        for ri, row in enumerate(table.rows):
            for ci, cell in enumerate(row.cells):
                if any(cell._tc is t for t in seen):
                    continue
                seen.append(cell._tc)
                out.append({"table": ti, "row": ri, "col": ci, "text": cell.text})
    return out


def _cells_as_text(cells: list[dict]) -> str:
    lines = []
    for c in cells:
        t = c["text"].replace("\n", "⏎")
        lines.append(f"table={c['table']} row={c['row']} col={c['col']}: {t!r}")
    return "\n".join(lines)


def fill_form_fields(llm, profile: MunicipalityProfile, colony: Colony, form_key: str, apply_date: dt.date | None) -> dict:
    form = profile.forms.get(form_key, {})
    template = profile.template_path(form_key)
    cells = template_cells(template) if template else []
    user = (
        f"様式：{form.get('label', form_key)}（テンプレートファイル：{form.get('template', '不明')}）\n"
        f"申請日：{apply_date.isoformat() if apply_date else '未定'}\n\n"
        "## コロニー情報\n" + colony.to_json()
    )
    if cells:
        user += "\n\n## 様式テンプレートのセル一覧（table/row/col と現在の文言）\n" + _cells_as_text(cells)
    out = llm.json(
        task=f"fill_{form_key}",
        system=FILL_SYSTEM,
        user=user,
        schema=FIELDS_SCHEMA,
        corpus=profile.corpus(),
        ctx={"profile": profile, "colony": colony, "form_key": form_key, "apply_date": apply_date, "cells": cells},
    )
    out.setdefault("cell_edits", [])
    return out


# ---- DOCX 生成 ----

def _set_cell_text(cell, text: str):
    """セルの文字を置き換える。最初の段落の書式（フォント）はできるだけ残す。"""
    paras = cell.paragraphs
    first = paras[0]
    font_size = None
    for r in first.runs:
        if r.font.size:
            font_size = r.font.size
            break
    for p in paras[1:]:
        p._element.getparent().remove(p._element)
    for r in list(first.runs):
        r._element.getparent().remove(r._element)
    lines = text.split("\n")
    run = first.add_run(lines[0])
    if font_size:
        run.font.size = font_size
    for line in lines[1:]:
        p = cell.add_paragraph()
        r = p.add_run(line)
        if font_size:
            r.font.size = font_size


def _apply_template(template_path: Path, edits: list[dict]) -> tuple[Document, list[dict]]:
    doc = Document(str(template_path))
    failed = []
    for e in edits:
        try:
            cell = doc.tables[e["table"]].cell(e["row"], e["col"])
            _set_cell_text(cell, e["text"])
        except Exception:
            failed.append(e)
    return doc, failed


def build_document(title: str, addressed_to: str, fields: list[dict], notes: list[str], footer: str = "") -> Document:
    """テンプレートが無い様式（第5号様式など）を python-docx でゼロから組む。"""
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(title, level=1)
    if addressed_to:
        doc.add_paragraph(addressed_to)
    doc.add_paragraph("")
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    for f in fields:
        row = table.add_row().cells
        row[0].text = f["label"]
        row[1].text = f["value"] or "　"
    if notes:
        doc.add_paragraph("")
        doc.add_paragraph("申請者への確認事項（提出前に削除してください）：")
        for n in notes:
            doc.add_paragraph(n, style="List Bullet")
    if footer:
        doc.add_paragraph("")
        doc.add_paragraph(footer)
    return doc


def render_docx(profile: MunicipalityProfile, form_key: str, filled: dict) -> tuple[bytes, str]:
    """(docx bytes, 生成方法ラベル) を返す。"""
    template = profile.template_path(form_key)
    if template and filled.get("cell_edits"):
        doc, failed = _apply_template(template, filled["cell_edits"])
        how = f"配布様式 {template.name} に転記（{len(filled['cell_edits']) - len(failed)}セル）"
        if failed:
            how += f"／転記できなかったセル {len(failed)}"
    else:
        doc = build_document(
            filled["form_title"], filled["addressed_to"], filled["fields"], filled["notes_for_applicant"],
            footer=f"※ {profile.name} 配布の様式テンプレートが無いため自動生成した代替です。提出時は配布様式に転記してください。",
        )
        how = "様式テンプレート未配置のため python-docx で生成"
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue(), how
