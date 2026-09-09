"""自治体プロファイルの読み込み。

municipality/<id>/profile.yaml と sources/ 内の一次資料（PDF/DOCX/MD）を
そのままテキスト化して LLM に渡す。規則の解釈はここでは一切しない。
"""
from __future__ import annotations

import dataclasses
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
MUNICIPALITY_DIR = ROOT / "municipality"


@dataclasses.dataclass
class SourceDoc:
    file: str
    role: str
    url: str
    text: str
    available: bool


@dataclasses.dataclass
class MunicipalityProfile:
    id: str
    name: str
    program_name: str
    office: str
    contacts: list[str]
    homepage: str
    sources: list[SourceDoc]
    forms: dict
    dir: Path

    def corpus(self) -> str:
        """LLM に渡す資料束。利用可能な資料だけを、役割ラベル付きで連結する。"""
        parts = []
        for s in self.sources:
            if not s.available:
                continue
            parts.append(f"<document file=\"{s.file}\" role=\"{s.role}\" url=\"{s.url}\">\n{s.text}\n</document>")
        return "\n\n".join(parts)

    def missing_sources(self) -> list[SourceDoc]:
        return [s for s in self.sources if not s.available]

    def template_path(self, form_key: str) -> Path | None:
        t = self.forms.get(form_key, {}).get("template")
        if not t:
            return None
        p = self.dir / "sources" / t
        return p if p.exists() else None


def _read_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _read_docx(path: Path) -> str:
    import docx

    d = docx.Document(str(path))
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            lines.append(" | ".join(cells))
    return "\n".join(lines)


def read_source_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".docx":
        return _read_docx(path)
    return path.read_text(encoding="utf-8")


def list_municipalities() -> list[str]:
    return sorted(p.name for p in MUNICIPALITY_DIR.iterdir() if (p / "profile.yaml").exists())


def load_profile(municipality_id: str = "kyoto") -> MunicipalityProfile:
    mdir = MUNICIPALITY_DIR / municipality_id
    meta = yaml.safe_load((mdir / "profile.yaml").read_text(encoding="utf-8"))
    sources: list[SourceDoc] = []
    for s in meta.get("sources", []):
        path = mdir / "sources" / s["file"]
        available = path.exists()
        text = ""
        if available:
            try:
                text = read_source_text(path)
                if len(text.strip()) < 50:  # 画像だけの PDF など。テキストが取れないので LLM には渡さない
                    available = False
                    text = ""
            except Exception as e:  # 読めない資料は「無い」扱いにして先へ進む
                text = ""
                available = False
                print(f"[profile] {path.name} を読めませんでした: {e}")
        sources.append(SourceDoc(file=s["file"], role=s.get("role", ""), url=s.get("url", ""), text=text, available=available))
    return MunicipalityProfile(
        id=meta["id"],
        name=meta["name"],
        program_name=meta.get("program_name", ""),
        office=meta.get("office", ""),
        contacts=meta.get("contacts", []),
        homepage=meta.get("homepage", ""),
        sources=sources,
        forms=meta.get("forms", {}),
        dir=mdir,
    )
